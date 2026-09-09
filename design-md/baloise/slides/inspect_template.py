#!/usr/bin/env python3
"""Inventory a corporate .potx/.pptx so its layouts can be reused rather than reinvented.

Reports, per file:
  - slide size in inches and points
  - theme fonts (major/minor) and the theme colour scheme
  - every slide master and slide layout, with each placeholder's type,
    index, name and geometry in points

Usage:  python3 inspect_template.py <file.potx|file.pptx> [--slides]
"""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU_PER_PT = 12700
EMU_PER_IN = 914400


def pt(v):
    return "-" if v is None else f"{Emu(v).emu / EMU_PER_PT:.0f}"


def describe_placeholder(ph):
    f = ph.placeholder_format
    try:
        ph_type = str(f.type).split(" ")[0]
    except ValueError:
        ph_type = "UNKNOWN"
    return (
        f"      idx={f.idx:<3} {ph_type:<18} "
        f"x={pt(ph.left):>5} y={pt(ph.top):>5} "
        f"w={pt(ph.width):>5} h={pt(ph.height):>5}  {ph.name}"
    )


def theme_info(prs):
    """Pull font scheme and colour scheme out of the first master's theme part."""
    from pptx.oxml.ns import qn

    master = prs.slide_masters[0]
    theme = master.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
    )
    root = theme._element if hasattr(theme, "_element") else None
    if root is None:
        from lxml import etree

        root = etree.fromstring(theme.blob)

    elements = root.find(qn("a:themeElements"))
    if elements is None:
        return

    scheme = elements.find(qn("a:fontScheme"))
    if scheme is not None:
        for role in ("a:majorFont", "a:minorFont"):
            node = scheme.find(qn(role))
            latin = node.find(qn("a:latin")) if node is not None else None
            if latin is not None:
                print(f"  {role.split(':')[1]:<10} {latin.get('typeface')}")

    colors = elements.find(qn("a:clrScheme"))
    if colors is not None:
        print("\nTheme colours")
        for child in colors:
            name = child.tag.split("}")[1]
            srgb = child.find(qn("a:srgbClr"))
            sys_clr = child.find(qn("a:sysClr"))
            if srgb is not None:
                value = "#" + srgb.get("val").upper()
            elif sys_clr is not None:
                value = "#" + (sys_clr.get("lastClr") or "?").upper()
            else:
                value = "?"
            print(f"  {name:<12} {value}")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    path = sys.argv[1]
    show_slides = "--slides" in sys.argv
    prs = Presentation(path)

    w, h = prs.slide_width, prs.slide_height
    print(f"=== {path} ===")
    print(
        f"Slide size: {w / EMU_PER_IN:.3f}in x {h / EMU_PER_IN:.3f}in "
        f"= {pt(w)}pt x {pt(h)}pt"
    )

    print("\nTheme fonts")
    theme_info(prs)

    for m, master in enumerate(prs.slide_masters):
        print(f"\n--- Master {m}: {master.name or '(unnamed)'} ---")
        for layout in master.slide_layouts:
            print(f"\n  Layout: {layout.name}")
            for ph in layout.placeholders:
                print(describe_placeholder(ph))

    if show_slides:
        print(f"\n--- {len(prs.slides)} existing slide(s) ---")
        for i, slide in enumerate(prs.slides, 1):
            print(f"\n  Slide {i}  (layout: {slide.slide_layout.name})")
            for shape in slide.shapes:
                kind = "PLACEHOLDER" if shape.is_placeholder else str(shape.shape_type)
                text = ""
                if shape.has_text_frame:
                    text = " | ".join(
                        p.text for p in shape.text_frame.paragraphs if p.text
                    )[:70]
                print(
                    f"    {kind:<20} x={pt(shape.left):>5} y={pt(shape.top):>5} "
                    f"w={pt(shape.width):>5} h={pt(shape.height):>5}  {text}"
                )


if __name__ == "__main__":
    main()
