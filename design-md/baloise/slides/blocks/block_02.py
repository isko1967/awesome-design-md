"""Block 2 - Persistent context & system instructions (slides 15-24).

Ten slides, full LEARN -> SEE -> TRY -> IMPROVE cycle.
Content is verbatim from the source deck; only the layout is ours.
"""
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from slidekit import (ACTION, ASIDE_W, ASIDE_X, BLUE, COL2_W, COL2_X, COL3_W,
                      COL3_X, CONTENT_TOP, CONTENT_W, GOOD, GREY_TEXT,
                      GUARDRAIL, MAIN_W, MARGIN, NEUTRAL, T_BODY, T_CARD,
                      T_HINT, T_LEAD, T_META, T_STATEMENT, badge, band,
                      blue_slide, card, divider, eyebrow, footer, lead,
                      listing, new_slide, runs, textbox, title, write)

LABEL = "02  Persistent context & system instructions"


def _foot(slide, phase, number):
    footer(slide, LABEL, phase, number)


def slide_15(prs):
    s = blue_slide(prs)
    divider(s, "02", ["Persistent context &", "system instructions"],
            "Stop repeating what stays the same.",
            chips=("LEARN", "SEE", "TRY", "IMPROVE"))
    _foot(s, None, 15)


def slide_16(prs):
    s = new_slide(prs)
    title(s, "Some instructions belong to the task. Others apply again and "
             "again.")
    lead(s, "Some information is needed for one task; some should shape many.")

    task = card(s, COL2_X[0], 150, COL2_W, 232, NEUTRAL.surface, pad=18)
    write(task, [
        ("Task-specific", T_CARD, True, GREY_TEXT, None),
        ("Summarise the main findings from this month’s claims report.",
         T_BODY, False, BLUE, 12),
        ("Changes every time.", T_HINT, False, GREY_TEXT, 12),
    ])
    persistent = card(s, COL2_X[1], 150, COL2_W, 232, GOOD.surface, pad=18)
    write(persistent, [
        ("Persistent", T_CARD, True, GOOD.accent, None),
        ("My primary audience is senior internal stakeholders.",
         T_BODY, False, BLUE, 12),
        ("Use concise business language.", T_BODY, False, BLUE, 4),
        ("Lead with the main conclusion.", T_BODY, False, BLUE, 4),
        ("Separate facts, assumptions and recommendations.",
         T_BODY, False, BLUE, 4),
        ("Applies repeatedly.", T_HINT, False, GOOD.accent, 10),
    ])

    band(s, "Key idea",
         "If an instruction stays true across many prompts, move it into "
         "persistent context.", GOOD)
    _foot(s, "LEARN", 16)


def slide_17(prs):
    s = new_slide(prs)
    title(s, "Save only the instructions that will still be true next time.")

    eyebrow(s, "Good candidates", w=MAIN_W)
    candidates = [
        ("Role & recurring work", "I work on cross-functional projects in "
         "insurance."),
        ("Typical audience", "I communicate with business owners, project "
         "teams and senior management."),
        ("Communication preferences", "Be concise, direct and structured."),
        ("Output defaults", "For recommendations, start with the "
         "recommendation, then rationale and next steps."),
        ("Recurring working rules", "Flag uncertainty. Do not invent missing "
         "information."),
    ]
    top, pitch = 146, 58
    for i, (labelt, example) in enumerate(candidates):
        y = top + i * pitch
        cell = card(s, MARGIN, y, MAIN_W, pitch - 8, NEUTRAL.surface, pad=10)
        runs(cell, [(f"{labelt}   ", True, BLUE), (example, False, BLUE)],
             size=T_HINT, anchor=MSO_ANCHOR.MIDDLE)

    # marginal material to the 203 aside at 13pt (overflow ladder, step 1)
    eyebrow(s, "Keep task-specific", x=ASIDE_X, w=ASIDE_W)
    aside = card(s, ASIDE_X, 146, ASIDE_W, 194, NEUTRAL.surface, pad=12)
    write(aside, [(t, T_HINT, False, BLUE, None if i == 0 else 8)
                  for i, t in enumerate((
                      "Current project facts",
                      "Temporary priorities",
                      "Individual customer information",
                      "One-off requirements"))],
          anchor=MSO_ANCHOR.MIDDLE)
    rel = card(s, ASIDE_X, 344, ASIDE_W, 96, NEUTRAL.surface, pad=12)
    write(rel, [
        ("ROLE PROMPTING", T_META, True, GREY_TEXT, None),
        ("Pair a role with concrete expectations, not just “act as an "
         "expert”.", T_HINT, False, BLUE, 6),
    ])
    _foot(s, "LEARN", 17)


def slide_18(prs):
    s = new_slide(prs)
    title(s, "Organise persistent instructions by what they control.")

    cats = [
        ("ROLE & CONTEXT", "What work do I repeatedly do?"),
        ("AUDIENCE", "Who usually receives my outputs?"),
        ("COMMUNICATION STYLE", "How should AI generally communicate?"),
        ("OUTPUT DEFAULTS", "Which structures do I repeatedly prefer?"),
        ("WORKING RULES", "What should AI do when information is missing?"),
    ]
    top, pitch = 146, 50
    for i, (labelt, q) in enumerate(cats):
        y = top + i * pitch
        cell = card(s, MARGIN, y, CONTENT_W, pitch - 8, NEUTRAL.surface,
                    pad=10)
        runs(cell, [(f"{labelt}    ", True, BLUE), (q, False, GREY_TEXT)],
             size=T_BODY, anchor=MSO_ANCHOR.MIDDLE)

    band(s, "Key idea",
         "Good persistent context is reusable enough to save effort, but "
         "specific enough to influence the result.", GOOD)
    _foot(s, "LEARN", 18)


def slide_19(prs):
    s = new_slide(prs)
    title(s, "Repeated output requirements do not need to be rewritten every "
             "time.")

    eyebrow(s, "Example — management decision note")
    prompt = card(s, MARGIN, 146, CONTENT_W, 262, NEUTRAL.surface, pad=20)
    write(prompt, [
        ("When I ask for a management decision note, use exactly these "
         "sections:", T_BODY, False, BLUE, None),
        ("1  Decision required — what needs to be decided?",
         T_BODY, True, BLUE, 8),
        ("2  Recommendation — what do you recommend?", T_BODY, True, BLUE, 4),
        ("3  Rationale — what are the strongest reasons?",
         T_BODY, True, BLUE, 4),
        ("4  Risks / trade-offs — what should decision-makers know?",
         T_BODY, True, BLUE, 4),
        ("5  Next step — what happens after the decision?",
         T_BODY, True, BLUE, 4),
        ("If essential information is missing, flag it rather than filling "
         "the gap yourself.", T_HINT, False, GREY_TEXT, 12),
    ])
    band(s, "Use when",
         "The same output type returns again and again.", GOOD)
    _foot(s, "LEARN", 19)


def slide_20(prs):
    s = new_slide(prs)
    title(s, "Persistent context only helps if you can keep it accurate.")

    eyebrow(s, "Hard to maintain", x=COL2_X[0], w=COL2_W)
    eyebrow(s, "Easier to maintain", x=COL2_X[1], w=COL2_W, colour=GOOD.accent)

    hard = card(s, COL2_X[0], 146, COL2_W, 250, NEUTRAL.surface, pad=20)
    write(hard, [
        ("I work in insurance and mostly do projects and normally I like "
         "short outputs although sometimes I need more detail and when I ask "
         "for recommendations…", T_BODY, False, BLUE, None),
    ], anchor=MSO_ANCHOR.MIDDLE)

    easy = card(s, COL2_X[1], 146, COL2_W, 250, GOOD.surface, pad=18)
    for i, (labelt, val) in enumerate([
            ("ROLE", "Cross-functional project work in insurance"),
            ("AUDIENCE", "Business owners · senior management · project teams"),
            ("STYLE", "Direct · concise · structured"),
            ("OUTPUT DEFAULT", "Recommendation → rationale → next step"),
            ("GUARDRAILS", "Flag assumptions · do not invent information")]):
        rows = [(labelt, T_META, True, GOOD.accent, None if i == 0 else 8),
                (val, T_HINT, False, BLUE, 2)]
        if i == 0:
            write(easy, rows)
        else:
            for text, size, bold, colour, before in rows:
                p = easy.text_frame.add_paragraph()
                p.alignment = PP_ALIGN.LEFT
                p.space_before = Pt(before or 0)
                p.line_spacing = Pt(round(size * 1.35))
                run = p.add_run()
                run.text = text
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = colour

    band(s, "Key idea",
         "Structure is not another technique. It makes persistent "
         "instructions easier to review and update.", GOOD)
    _foot(s, "LEARN", 20)


def slide_21(prs):
    s = new_slide(prs)
    title(s, "How much can the same short prompt improve when the context is "
             "already there?")

    eyebrow(s, "Scenario — claims operations notes", w=MAIN_W)
    notes = card(s, MARGIN, 146, MAIN_W, 200, NEUTRAL.surface, pad=16)
    write(notes, [(t, T_BODY, False, BLUE, None if i == 0 else 6)
                  for i, t in enumerate((
                      "Claims volume increased after severe weather.",
                      "Average processing time increased.",
                      "Additional external capacity activated.",
                      "Customer communication unchanged.",
                      "Management wants a weekly update."))],
          anchor=MSO_ANCHOR.MIDDLE)

    eyebrow(s, "Prompt — both times", x=ASIDE_X, w=ASIDE_W)
    prompt = card(s, ASIDE_X, 146, ASIDE_W, 92, GOOD.surface, pad=12)
    write(prompt, [("Draft a weekly update based on these notes.",
                    T_BODY, True, BLUE, None)], anchor=MSO_ANCHOR.MIDDLE)
    runs_box = card(s, ASIDE_X, 254, ASIDE_W, 92, NEUTRAL.surface, pad=12)
    write(runs_box, [
        ("We run it twice", T_META, True, GREY_TEXT, None),
        ("1  Without context", T_HINT, True, BLUE, 8),
        ("2  With context", T_HINT, True, BLUE, 4),
    ], anchor=MSO_ANCHOR.MIDDLE)

    band(s, "Watch for",
         "Relevance · prioritisation · structure · audience fit · repeated "
         "instructions")
    _foot(s, "SEE", 21)


def slide_22(prs):
    s = new_slide(prs)
    title(s, "Same prompt. Different context.")

    eyebrow(s, "Without persistent context", x=COL2_X[0], w=COL2_W)
    eyebrow(s, "With persistent context", x=COL2_X[1], w=COL2_W,
            colour=GOOD.accent)

    before = card(s, COL2_X[0], 146, COL2_W, 296, NEUTRAL.surface, pad=20)
    write(before, [
        ("Claims volumes increased following recent severe weather events. "
         "Processing times have also increased, and the team has taken "
         "several measures to address the situation. Additional external "
         "capacity has been activated.", T_BODY, False, BLUE, None),
    ], anchor=MSO_ANCHOR.MIDDLE)

    after = card(s, COL2_X[1], 146, COL2_W, 296, GOOD.surface, pad=20)
    rows = []
    for i, (heading, text) in enumerate([
            ("Status — Attention required",
             "Severe-weather claims have increased volumes and processing "
             "times."),
            ("Response", "Additional external capacity has been activated."),
            ("Customer impact",
             "No communication change is required at this stage."),
            ("Management watchpoint",
             "Monitor processing times over the coming week.")]):
        rows.append((heading, T_BODY, True, GOOD.accent, None if i == 0 else 12))
        rows.append((text, T_BODY, False, BLUE, 0))
    write(after, rows, anchor=MSO_ANCHOR.MIDDLE)

    band(s, "What changed?",
         "Prioritisation · structure · audience fit", GOOD)
    _foot(s, "SEE", 22)


def slide_23(prs):
    s = new_slide(prs)
    title(s, "Build context you would actually reuse.", width=729)
    lead(s, "Your task")
    badge(s, "TIME  5 MIN")

    steps = [
        ("1", "Think of work you repeatedly use AI for.",
         "Management updates, meeting prep, project communication, analysis, "
         "decision papers."),
        ("2", "Identify what remains stable.",
         "Role, typical audience, communication style, preferred structure, "
         "recurring rules."),
        ("3", "Write 4–6 persistent instructions.", None),
        ("4", "Test them with one short task prompt.",
         "For example: draft a project update from these notes."),
    ]
    top, pitch = 146, 74
    for i, (num, text, hint) in enumerate(steps):
        y = top + i * pitch
        cell = card(s, MARGIN, y, CONTENT_W, pitch - 10, ACTION.surface,
                    pad=12)
        rows = [(f"{num}   ", True, ACTION.accent), (text, True, BLUE)]
        runs(cell, rows, size=T_BODY, anchor=MSO_ANCHOR.MIDDLE
             if not hint else MSO_ANCHOR.TOP)
        if hint:
            p = cell.text_frame.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(2)
            p.line_spacing = Pt(18)
            run = p.add_run()
            run.text = hint
            run.font.size = Pt(T_HINT)
            run.font.color.rgb = GREY_TEXT

    band(s, "Final question",
         "Which instructions did you no longer need to repeat?", ACTION)
    _foot(s, "TRY", 23)


def slide_24(prs):
    s = new_slide(prs)
    title(s, "What became reusable?")
    lead(s, "Which type of information did you move out of the individual "
            "prompt?")

    x = MARGIN
    for text in ("Role / context", "Audience", "Style", "Output format",
                 "Working rules"):
        chip = card(s, x, 168, 163, 96, None, outline=ACTION.accent)
        write(chip, [(text, T_LEAD, False, BLUE, None)],
              anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        x += 163 + 16

    cta = card(s, MARGIN, 300, CONTENT_W, 132, ACTION.surface, pad=24)
    write(cta, [
        ("In the chat", T_BODY, True, ACTION.accent, None),
        ("Share one instruction you would want AI to remember across tasks.",
         T_STATEMENT, False, BLUE, 8),
    ], anchor=MSO_ANCHOR.MIDDLE)
    _foot(s, "IMPROVE", 24)


SLIDES = [slide_15, slide_16, slide_17, slide_18, slide_19, slide_20,
          slide_21, slide_22, slide_23, slide_24]
