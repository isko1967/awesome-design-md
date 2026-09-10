"""Block 7 - Transfer & close (slides 50-54).

The practice that protects transfer time, the make-it-reusable list, the
five prompt-library fields, the decision matrix mapping problems to
techniques, and the tangerine close. Resource links stay as placeholders.
Content is verbatim from the source deck; only the layout is ours.
"""
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from slidekit import (task_header, ACTION, BLUE, COL2_W, COL2_X, CONTENT_W, GOOD, GREEN,
                      GREY_TEXT, GUARDRAIL, MARGIN, NEUTRAL, T_BODY, T_CARD,
                      T_DIVIDER, T_HINT, T_LEAD, T_META, T_NUMBER, T_STATEMENT,
                      badge, band, blue_slide, card, divider, eyebrow, footer,
                      lead, listing, new_slide, runs, tangerine_slide,
                      textbox, title, write)

LABEL = "07  Transfer & close"


def _foot(slide, phase, number):
    footer(slide, LABEL, phase, number)


def slide_50(prs):
    s = new_slide(prs)
    title(s, "Build the prompt you will actually use next.", width=729)
    task_header(s, "Your task", "7 min")

    steps = [
        ("1", "Choose a recurring task.",
         "Pick something you expect to do again."),
        ("2", "Write the complete prompt.",
         "Not notes — the version you would actually use."),
        ("3", "Apply at least one technique from today.",
         "Few-shot? Persistent context? Decomposition? Self-critique? "
         "Debugging?"),
        ("4", "Test and adjust it.", None),
        ("5", "Save it where you will find it again.", None),
    ]
    top, pitch = 146, 62
    for i, (num, text, hint) in enumerate(steps):
        y = top + i * pitch
        cell = card(s, MARGIN, y, CONTENT_W, pitch - 8, ACTION.surface,
                    pad=8)
        runs(cell, [(f"{num}   ", True, ACTION.accent), (text, True, BLUE)],
             size=T_BODY, anchor=MSO_ANCHOR.TOP if hint else MSO_ANCHOR.MIDDLE)
        if hint:
            p = cell.text_frame.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(2)
            p.line_spacing = Pt(16)
            r = p.add_run()
            r.text = hint
            r.font.size = Pt(T_HINT)
            r.font.color.rgb = GREY_TEXT
    _foot(s, "Transfer", 50)


def slide_51(prs):
    s = new_slide(prs)
    title(s, "Turn a good prompt into a reusable prompt.")
    task_header(s, "Before you save it")

    items = [
        ("Name it", "Give it a clear, task-based name."),
        ("Generalise it", "Replace one-off details with reusable "
         "placeholders."),
        ("Explain when to use it", "Add one short sentence."),
        ("Clean it", "Remove information that should not be reused."),
        ("Capture the learning", "What made the prompt work?"),
    ]
    top, pitch = 150, 48
    for i, (head, body) in enumerate(items):
        y = top + i * pitch
        cell = card(s, MARGIN, y, CONTENT_W, pitch - 8, NEUTRAL.surface,
                    pad=10)
        runs(cell, [(head + "   ", True, GREEN), (body, False, BLUE)],
             size=T_BODY, anchor=MSO_ANCHOR.MIDDLE)

    band(s, "Then", "Add it to the Prompt Library.", GOOD)
    _foot(s, "Transfer", 51)


def slide_52(prs):
    s = new_slide(prs)
    title(s, "What makes a prompt reusable?")
    task_header(s, "Capture five things")

    fields = [
        ("1", "Situation", "When is this useful?"),
        ("2", "Use case", "What does it help you do?"),
        ("3", "Prompt", "What should the next user copy and adapt?"),
        ("4", "Technique", "What improved the result?"),
        ("5", "Insight", "What should the next user know?"),
    ]
    top, pitch = 148, 46
    for i, (num, head, q) in enumerate(fields):
        y = top + i * pitch
        role = GOOD if i == 4 else NEUTRAL
        cell = card(s, MARGIN, y, CONTENT_W, pitch - 8, role.surface, pad=10)
        runs(cell, [(f"{num}   ", True,
                     GOOD.accent if i == 4 else GREEN),
                    (f"{head}   ", True, BLUE), (q, False, GREY_TEXT)],
             size=T_BODY, anchor=MSO_ANCHOR.MIDDLE)

    band(s, "Why", "The insight makes the prompt reusable.", GOOD)
    _foot(s, "Transfer", 52)


def slide_53(prs):
    """The decision matrix - five problems mapped to five techniques."""
    s = new_slide(prs)
    title(s, "Five prompting problems. Five techniques.")

    rows = [
        ("The target is difficult to describe.",
         "01 Few-shot & negative examples", "Show instead of describe."),
        ("I keep repeating the same instructions.",
         "02 Persistent context", "Move stable instructions out."),
        ("The task is complex or multi-step.",
         "03 Decomposition", "Check the approach before executing."),
        ("I need better quality control.",
         "04 Self-critique", "Define criteria and check against them."),
        ("The result failed and I don’t know why.",
         "05 Debugging", "Diagnose instead of starting over."),
    ]
    top, pitch = 146, 50
    for i, (problem, technique, answer) in enumerate(rows):
        y = top + i * pitch
        prob = card(s, MARGIN, y, 300, pitch - 8, NEUTRAL.surface, pad=10)
        write(prob, [(problem, T_HINT, False, BLUE, None)],
              anchor=MSO_ANCHOR.MIDDLE)
        tech = card(s, MARGIN + 312, y, CONTENT_W - 312, pitch - 8,
                    GOOD.surface, pad=10)
        runs(tech, [(technique + "   ", True, GOOD.accent),
                    (answer, False, BLUE)], size=T_HINT,
             anchor=MSO_ANCHOR.MIDDLE)

    band(s, "Choose",
         "Pick the technique for the problem — not because it sounds "
         "advanced.", GOOD)
    _foot(s, None, 53)


def slide_54(prs):
    """The close - a full-bleed tangerine landmark with resource links."""
    s = tangerine_slide(prs)
    write(textbox(s, MARGIN, 108, CONTENT_W, 70),
          [("Keep practising.", T_DIVIDER, True, BLUE, None)])

    cards = [
        ("Academy", "[ACADEMY LINK]",
         "Refresh the fundamentals and revisit today’s techniques."),
        ("Prompt Library", "[PROMPT LIBRARY LINK]",
         "Reuse proven prompts and contribute your own."),
        ("Questions", "[PROGRAMME CONTACT / MAILBOX]",
         "Ask the programme team."),
    ]
    w = (CONTENT_W - 2 * 20) / 3
    x = MARGIN
    for head, link, body in cards:
        cell = card(s, x, 190, w, 150, NEUTRAL.surface, pad=16)
        write(cell, [
            (head, T_CARD, True, BLUE, None),
            (link, T_HINT, True, ACTION.accent, 10),
            (body, T_HINT, False, GREY_TEXT, 10),
        ], anchor=MSO_ANCHOR.TOP)
        x += w + 20

    close = card(s, MARGIN, 360, CONTENT_W, 62, ACTION.surface, pad=16)
    write(close, [("Save one prompt. Reuse it. Improve it.",
                   T_STATEMENT, True, BLUE, None)],
          anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    footer(s, LABEL, None, 54)


SLIDES = [slide_50, slide_51, slide_52, slide_53, slide_54]
