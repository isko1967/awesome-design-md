"""Block 1 - Few-shot & negative examples (slides 7-14).

Content is verbatim from the source deck; only the layout is ours.
"""
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from slidekit import (CHOICE, ACTION, AVOID, COL4_W, COL4_X, blue_slide, divider,
                      BAND_TOP, BLUE, COL2_W, COL2_X, COL3_W,
                      COL3_X, CONTENT_TOP, CONTENT_W, GOOD, GREY_TEXT,
                      GUARDRAIL, MARGIN,
                      NEUTRAL, T_BODY, T_CARD, T_DIVIDER, T_HINT, T_LEAD,
                      T_META, T_NUMBER, T_STATEMENT, WHITE, badge, band, card,
                      eyebrow, footer, lead, new_slide, runs, textbox, title,
                      write)

LABEL = "01  Few-shot & negative examples"


def _foot(slide, phase, number):
    footer(slide, LABEL, phase, number)


def slide_07(prs):
    s = blue_slide(prs)
    divider(s, "01", ["Few-shot &", "negative examples"],
            "Show the model what good looks like.",
            chips=("LEARN", "SEE", "TRY", "IMPROVE"))
    _foot(s, None, 7)


def slide_08(prs):
    s = new_slide(prs)
    title(s, "Few-shot prompting gives the model examples to learn from.")
    lead(s, "Instead of only describing what you want, provide one or more "
            "examples of the desired output.")

    eyebrow(s, "What the model infers")
    for x, pattern in zip(COL4_X, ("tone and level of formality",
                                   "structure and formatting",
                                   "length and level of detail",
                                   "wording and style")):
        tile = card(s, x, 142, COL4_W, 64, NEUTRAL.surface, pad=14)
        write(tile, [(pattern, 14, True, BLUE, None)],
              anchor=MSO_ANCHOR.MIDDLE)

    eyebrow(s, "Example", y=230)
    plain = card(s, COL2_X[0], 252, COL2_W, 164, NEUTRAL.surface)
    write(plain, [
        ("Instruction only", T_LEAD, True, GREY_TEXT, None),
        ("Summarise this claims update for senior management.",
         14, False, BLUE, 10),
    ])
    guided = card(s, COL2_X[1], 252, COL2_W, 164, GOOD.surface)
    write(guided, [
        ("With an example", T_LEAD, True, GOOD.accent, None),
        ("Summarise this claims update for senior management.",
         14, False, BLUE, 10),
        ("Use the example below as a reference for tone, structure and level "
         "of detail.", 14, False, BLUE, 6),
        ("Example: [insert strong previous update]", 14, True, BLUE, 8),
    ])

    related = textbox(s, MARGIN, 424, CONTENT_W, 20)
    runs(related, [("Related concept   ", True, GREY_TEXT),
                   ("Zero-shot means instructions without examples. Few-shot "
                    "means one or more examples are provided as guidance.",
                    False, BLUE)], size=T_HINT, anchor=MSO_ANCHOR.TOP)

    band(s, "Key idea",
         "Examples reduce how much the model has to interpret.")
    _foot(s, "LEARN", 8)


def slide_09(prs):
    s = new_slide(prs)
    title(s, "Negative examples show the model what to avoid.")
    lead(s, "Imagine you regularly create management updates about claims "
            "performance.")

    for x, question, answer in (
            (COL2_X[0], "A positive example answers",
             "What should the output look like?"),
            (COL2_X[1], "A negative example answers",
             "What should the output not look like?")):
        ask = card(s, x, CONTENT_TOP, COL2_W, 84, NEUTRAL.surface)
        write(ask, [(question, T_BODY, False, GREY_TEXT, None),
                    (answer, T_LEAD, True, BLUE, 4)])

    positive = card(s, COL2_X[0], 220, COL2_W, 216, GOOD.surface)
    write(positive, [
        ("Positive example", T_CARD, True, GOOD.accent, None),
        ("Status: Claims volumes increased by 12%.", T_BODY, False, BLUE, 12),
        ("Impact: Processing time is temporarily above target.",
         T_BODY, False, BLUE, 2),
        ("Action: Additional capacity has been activated.",
         T_BODY, False, BLUE, 2),
        ("→ gives direction", T_BODY, True, GOOD.accent, 14),
    ])
    negative = card(s, COL2_X[1], 220, COL2_W, 216, AVOID.surface)
    write(negative, [
        ("Negative example", T_CARD, True, AVOID.accent, None),
        ("We would like to provide an update regarding various developments "
         "that have taken place within claims operations over the past "
         "reporting period…", T_BODY, False, BLUE, 12),
        ("→ sets a boundary", T_BODY, True, AVOID.accent, 14),
    ])

    band(s, "Key idea",
         "You do not always need both. Add a negative example when avoiding "
         "a specific pattern matters.")
    _foot(s, "LEARN", 9)


def slide_10(prs):
    s = new_slide(prs)
    title(s, "Tell the model what to learn from each example.")

    eyebrow(s, "What to specify", w=COL2_W)
    positive = card(s, MARGIN, 146, COL2_W, 142, GOOD.surface)
    write(positive, [
        ("Positive example", T_CARD, True, GOOD.accent, None),
        ("Follow its tone, structure and level of detail.",
         T_BODY, False, BLUE, 10),
    ], anchor=MSO_ANCHOR.MIDDLE)
    negative = card(s, MARGIN, 300, COL2_W, 142, AVOID.surface)
    write(negative, [
        ("Negative example", T_CARD, True, AVOID.accent, None),
        ("Avoid its unnecessary background, vague wording and lack of clear "
         "actions.", T_BODY, False, BLUE, 10),
    ], anchor=MSO_ANCHOR.MIDDLE)

    eyebrow(s, "The prompt", x=COL2_X[1], w=COL2_W)
    prompt = card(s, COL2_X[1], 146, COL2_W, 296, NEUTRAL.surface)
    write(prompt, [
        ("Use the examples below as guidance for the new output.",
         T_BODY, False, BLUE, None),
        ("Positive example:", T_BODY, True, BLUE, 12),
        ("[insert example]", T_BODY, False, GREY_TEXT, 0),
        ("Negative example:", T_BODY, True, BLUE, 12),
        ("[insert example]", T_BODY, False, GREY_TEXT, 0),
        ("Now complete this task:", T_BODY, True, BLUE, 12),
        ("[insert task]", T_BODY, False, GREY_TEXT, 0),
        ("Do not copy the content of the examples. Use them only as guidance "
         "for how the new output should be written.",
         T_HINT, False, GREY_TEXT, 14),
    ])

    band(s, "Use when",
         "The desired quality is easier to demonstrate than to describe.")
    _foot(s, "LEARN", 10)


def slide_11(prs):
    s = new_slide(prs)
    title(s, "Can examples make a vague instruction more precise?")
    lead(s, "Severe weather has increased claims volumes. Management needs a "
            "short update.")

    eyebrow(s, "Starting prompt")
    start = card(s, MARGIN, 146, CONTENT_W, 62, NEUTRAL.surface)
    write(start, [("Turn these notes into a short update for senior "
                   "management.", T_LEAD, False, BLUE, None)],
          anchor=MSO_ANCHOR.MIDDLE)

    eyebrow(s, "Version 2 adds", x=COL2_X[0], w=COL2_W, y=226)
    eyebrow(s, "Watch for", x=COL2_X[1], w=COL2_W, y=226)

    for y, kind, purpose, role in (
            (252, "One positive example", "to show the desired style", GOOD),
            (366, "One negative example", "to show what to avoid", AVOID)):
        tile = card(s, COL2_X[0], y, COL2_W, 100, role.surface, pad=18)
        write(tile, [(kind, T_CARD, True, role.accent, None),
                     (purpose, T_BODY, False, BLUE, 4)],
              anchor=MSO_ANCHOR.MIDDLE)

    watch = card(s, COL2_X[1], 252, COL2_W, 214, NEUTRAL.surface, pad=18)
    write(watch, [(item, T_BODY, False, BLUE, None if i == 0 else 6)
                  for i, item in enumerate(
                      ("prioritisation", "structure", "level of detail",
                       "tone", "action orientation"))],
          anchor=MSO_ANCHOR.MIDDLE)
    _foot(s, "SEE", 11)


def slide_12(prs):
    s = new_slide(prs)
    title(s, "Same task. Clearer guidance.")

    eyebrow(s, "Without examples", x=COL2_X[0], w=COL2_W)
    eyebrow(s, "With examples", x=COL2_X[1], w=COL2_W, colour=GOOD.accent)

    before = card(s, COL2_X[0], 146, COL2_W, 296, NEUTRAL.surface, pad=20)
    write(before, [
        ("Recent severe weather events have led to an increase in claims "
         "volumes. The claims organisation is currently taking several "
         "actions to manage the situation. Additional capacity has been made "
         "available, and developments will continue to be monitored.",
         T_BODY, False, BLUE, None),
    ], anchor=MSO_ANCHOR.MIDDLE)

    after = card(s, COL2_X[1], 146, COL2_W, 296, GOOD.surface, pad=20)
    rows = []
    for i, (heading, text) in enumerate([
            ("Status — Attention required",
             "Severe-weather claims increased volumes and processing times."),
            ("Response", "Additional external capacity has been activated."),
            ("Customer impact",
             "No change to current customer communication is required."),
            ("Next watchpoint",
             "Reassess processing times at the end of the week.")]):
        rows.append((heading, T_BODY, True, GOOD.accent, None if i == 0 else 12))
        rows.append((text, T_BODY, False, BLUE, 0))
    write(after, rows, anchor=MSO_ANCHOR.MIDDLE)

    band(s, "What changed?",
         "More specific · more structured · more actionable", GOOD)  # verbatim
    _foot(s, "SEE", 12)


def slide_13(prs):
    s = new_slide(prs)
    title(s, "Use examples to guide your own output.", width=729)
    lead(s, "Your task")
    badge(s, "Time  5 min")

    cells = [
        ("1", "Choose a recurring text-based task.",
         "For example a stakeholder update, meeting summary, decision note "
         "or internal announcement.", ACTION),
        ("2", "Add one positive example.",
         "Choose something that represents the quality you want.", ACTION),
        ("3", "Optional: add one negative example.",
         "Use it if there is a recurring pattern you want to avoid.", ACTION),
        ("4", "Tell AI what to learn from the examples.",
         "Tone? Structure? Length? Level of detail?", ACTION),
        ("5", "Run the prompt and compare the result.", None, ACTION),
        (None, "Before you start",
         "Keep all information generic. Do not enter personal, customer or "
         "confidential information.", GUARDRAIL),
    ]
    for i, (number, text, hint, role) in enumerate(cells):
        x = COL2_X[i % 2]
        y = 158 + (i // 2) * 108
        cell = card(s, x, y, COL2_W, 100, role.surface, pad=16)
        head = [(f"{number}   ", True, role.accent)] if number else []
        runs(cell, head + [(text, True, BLUE)], size=T_BODY,
             anchor=MSO_ANCHOR.TOP)
        if hint:
            para = cell.text_frame.add_paragraph()
            para.alignment = PP_ALIGN.LEFT
            para.space_before = Pt(4)
            para.line_spacing = Pt(18)
            run = para.add_run()
            run.text = hint
            run.font.size = Pt(T_HINT)
            run.font.color.rgb = GREY_TEXT if role is ACTION else BLUE
    _foot(s, "TRY", 13)


def slide_14(prs):
    s = new_slide(prs)
    title(s, "Did the examples give you more control?")
    lead(s, "Compare the result with what you would normally receive.")

    write(textbox(s, MARGIN, 140, 500, 28),
          [("What changed most?", T_CARD, True, BLUE, None)])

    x = MARGIN
    for text in ("Tone", "Structure", "Level of detail", "Wording",
                 "Nothing meaningful"):
        chip = card(s, x, 180, 163, 110, None,
                    outline=CHOICE.surface)
        write(chip, [(text, T_LEAD, False, WHITE, None)],
              anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        x += 163 + 16

    cta = card(s, MARGIN, 326, CONTENT_W, 124, ACTION.surface, pad=24)
    write(cta, [
        ("In the chat", T_BODY, True, ACTION.accent, None),
        ("Share the one thing that changed most.", T_STATEMENT, False, BLUE, 8),
    ], anchor=MSO_ANCHOR.MIDDLE)
    _foot(s, "IMPROVE", 14)


SLIDES = [slide_07, slide_08, slide_09, slide_10, slide_11, slide_12,
          slide_13, slide_14]
